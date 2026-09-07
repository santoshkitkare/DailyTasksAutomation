"""Contacts workbook parsing and per-row validation (PRD section 43).

The governing rule: one bad row must never cost the other ninety-nine their
greetings. Every row is validated independently and failures are collected and
reported, not raised.
"""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import load_workbook

from ..contracts import ContactRecord, EventType, InvalidContactRow, Relationship
from ..occasion_dates import UNKNOWN_YEAR

logger = logging.getLogger(__name__)

REQUIRED_COLUMNS = ("EventType", "EventDate", "FullName", "EmailAddress")
OPTIONAL_COLUMNS = ("MobileNumber", "Relationship")

# Accepted written date formats, tried in order. Day-first before month-first
# because the contact list is maintained in India; an unambiguous ISO date is
# tried first so it can never be misread.
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d-%m-%Y",
    "%d/%m/%Y",
    "%d %B %Y",
    "%d %b %Y",
    "%B %d, %Y",
    "%b %d, %Y",
    "%d-%m",
    "%d/%m",
    "%d %B",
    "%d %b",
)

# Deliberately permissive: the goal is to catch typos and empty cells, not to
# re-implement RFC 5322. A wrong-but-plausible address fails at send time and
# is reported then.
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[A-Za-z]{2,}$")


class WorkbookError(RuntimeError):
    """The workbook itself is unusable - wrong sheet, missing columns."""


@dataclass
class ContactsParseResult:
    contacts: list[ContactRecord]
    invalid: list[InvalidContactRow]
    duplicates: list[str]

    @property
    def total_rows(self) -> int:
        return len(self.contacts) + len(self.invalid)

    def summary(self) -> str:
        parts = [f"{self.total_rows} records", f"{len(self.contacts)} valid"]
        if self.invalid:
            parts.append(f"{len(self.invalid)} invalid")
        if self.duplicates:
            parts.append(f"{len(self.duplicates)} duplicate")
        return ", ".join(parts)


def parse_contacts(data: bytes, worksheet: str) -> ContactsParseResult:
    """Parse the workbook into validated contacts plus a report of failures."""
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        if worksheet not in workbook.sheetnames:
            raise WorkbookError(
                f"Worksheet {worksheet!r} not found. The workbook contains: "
                f"{', '.join(workbook.sheetnames)}. Fix occasion.worksheet in "
                "config.yaml."
            )
        sheet = workbook[worksheet]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise WorkbookError(f"Worksheet {worksheet!r} is empty.")

    header_map = _map_headers(rows[0])
    missing = [c for c in REQUIRED_COLUMNS if c not in header_map]
    if missing:
        raise WorkbookError(
            f"Worksheet {worksheet!r} is missing required column(s): "
            f"{', '.join(missing)}. Required columns are "
            f"{', '.join(REQUIRED_COLUMNS)}."
        )

    contacts: list[ContactRecord] = []
    invalid: list[InvalidContactRow] = []
    duplicates: list[str] = []
    seen: set[tuple[str, str]] = set()

    for offset, row in enumerate(rows[1:], start=2):
        raw = {name: row[index] for name, index in header_map.items() if index < len(row)}
        if all(value in (None, "") for value in raw.values()):
            continue  # blank spacer row

        try:
            contact = _build_contact(offset, raw)
        except ValueError as exc:
            invalid.append(
                InvalidContactRow(
                    row_number=offset, reason=str(exc), raw=_stringify(raw)
                )
            )
            continue

        identity = (str(contact.event_type), contact.email_address.lower())
        if identity in seen:
            duplicates.append(
                f"Row {offset}: duplicate {contact.event_type} entry for "
                f"{contact.email_address}"
            )
            continue
        seen.add(identity)
        contacts.append(contact)

    return ContactsParseResult(contacts=contacts, invalid=invalid, duplicates=duplicates)


def _map_headers(header_row: tuple) -> dict[str, int]:
    """Match headers case- and space-insensitively so 'Full Name' works."""
    wanted = {
        _normalise_header(name): name for name in REQUIRED_COLUMNS + OPTIONAL_COLUMNS
    }
    mapping: dict[str, int] = {}
    for index, cell in enumerate(header_row):
        if cell is None:
            continue
        key = _normalise_header(str(cell))
        if key in wanted and wanted[key] not in mapping:
            mapping[wanted[key]] = index
    return mapping


def _normalise_header(value: str) -> str:
    return re.sub(r"[\s_-]+", "", value).casefold()


def _build_contact(row_number: int, raw: dict) -> ContactRecord:
    event_type = _parse_event_type(raw.get("EventType"))
    event_date = _parse_event_date(raw.get("EventDate"))

    full_name = _clean(raw.get("FullName"))
    if not full_name:
        raise ValueError("FullName is empty")

    email = _clean(raw.get("EmailAddress")).lower()
    if not email:
        raise ValueError("EmailAddress is empty")
    if not EMAIL_RE.match(email):
        raise ValueError(f"EmailAddress {email!r} is not a valid address")

    return ContactRecord(
        row_number=row_number,
        event_type=event_type,
        event_date=event_date,
        full_name=full_name,
        email_address=email,
        mobile_number=_clean(raw.get("MobileNumber")),
        relationship=Relationship.parse(_clean(raw.get("Relationship"))),
    )


def _parse_event_type(value) -> EventType:  # noqa: ANN001 - openpyxl cell value
    text = _clean(value)
    if not text:
        raise ValueError("EventType is empty")
    normalised = text.casefold().rstrip("s")
    for member in EventType:
        if member.value.casefold().rstrip("s") == normalised:
            return member
    supported = ", ".join(m.value for m in EventType)
    raise ValueError(f"EventType {text!r} is not supported (expected one of: {supported})")


def _parse_event_date(value) -> date:  # noqa: ANN001 - openpyxl cell value
    """Accept a real Excel date or a written one.

    A date with no year (for example "7 September") is stored against year 1904
    as a sentinel: only month and day are ever used for matching, and 1904 is a
    leap year so 29 February survives the round trip.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = _clean(value)
    if not text:
        raise ValueError("EventDate is empty")

    for fmt in DATE_FORMATS:
        # A year-less format is parsed against UNKNOWN_YEAR explicitly rather
        # than relying on strptime's default. The default is 1900, which is not
        # a leap year, so "29 February" would fail to parse at all - and CPython
        # is deprecating the year-less default outright.
        yearless = "%Y" not in fmt
        candidate = f"{text} {UNKNOWN_YEAR}" if yearless else text
        pattern = f"{fmt} %Y" if yearless else fmt
        try:
            parsed = datetime.strptime(candidate, pattern)
        except ValueError:
            continue
        return parsed.date()

    raise ValueError(f"EventDate {text!r} could not be read as a date")


def _clean(value) -> str:  # noqa: ANN001 - openpyxl cell value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _stringify(raw: dict) -> dict:
    return {key: _clean(value) for key, value in raw.items()}
